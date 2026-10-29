"""Build a research-update deck from the completed face/eye SAE experiments."""
import io
import sys
from pathlib import Path

PPTX_DEPS = Path("/tmp/pptx_runtime")
if PPTX_DEPS.is_dir():
    sys.path.insert(0, str(PPTX_DEPS))

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "sae-training/presentations/face_sae_research_update.pptx"
W, H = 13.333, 7.5
NAVY = RGBColor(23, 38, 58)
INK = RGBColor(40, 50, 64)
MUTED = RGBColor(102, 113, 127)
TEAL = RGBColor(0, 143, 147)
PALE = RGBColor(232, 244, 243)
PAPER = RGBColor(249, 250, 252)
WHITE = RGBColor(255, 255, 255)
ORANGE = RGBColor(220, 119, 76)

prs = Presentation()
prs.slide_width = Inches(W)
prs.slide_height = Inches(H)
prs.core_properties.title = "Localizing face-related activations in a pretrained SAE"
prs.core_properties.subject = "Face and eye SAE fine-tuning and held-out feature localization"
prs.core_properties.author = "Research update"
blank = prs.slide_layouts[6]


def textbox(slide, text, x, y, w, h, size=18, color=INK, bold=False,
            align=PP_ALIGN.LEFT, font="Aptos"):
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = shape.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.margin_left = Inches(.02)
    frame.margin_right = Inches(.02)
    frame.margin_top = Inches(.01)
    frame.margin_bottom = Inches(.01)
    paragraph = frame.paragraphs[0]
    paragraph.text = text
    paragraph.alignment = align
    paragraph.font.name = font
    paragraph.font.size = Pt(size)
    paragraph.font.bold = bold
    paragraph.font.color.rgb = color
    return shape


def notes(slide, text):
    slide.notes_slide.notes_text_frame.text = text


def new_slide(title, subtitle="", note=""):
    slide = prs.slides.add_slide(blank)
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = PAPER
    band = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(.16), Inches(H))
    band.fill.solid()
    band.fill.fore_color.rgb = TEAL
    band.line.fill.background()
    textbox(slide, title, .62, .34, 12.0, .58, 27, NAVY, True)
    if subtitle:
        textbox(slide, subtitle, .65, .98, 12.0, .38, 13, MUTED)
    rule = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(.65), Inches(1.48), Inches(12.0), Inches(.018))
    rule.fill.solid()
    rule.fill.fore_color.rgb = RGBColor(210, 220, 228)
    rule.line.fill.background()
    textbox(slide, f"{len(prs.slides):02d}", 12.2, 7.12, .35, .2, 9, MUTED,
            align=PP_ALIGN.RIGHT)
    if note:
        notes(slide, note)
    return slide


def card(slide, x, y, w, h, title, body, fill=WHITE, accent=TEAL,
         title_size=16, body_size=14):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.color.rgb = RGBColor(222, 230, 236)
    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(.07), Inches(h))
    bar.fill.solid()
    bar.fill.fore_color.rgb = accent
    bar.line.fill.background()
    if h < 1.0:
        textbox(slide, title, x + .22, y + .08, w - .4, .24, title_size - 1, NAVY, True)
        textbox(slide, body, x + .22, y + .37, w - .4, h - .42, body_size, INK)
    else:
        textbox(slide, title, x + .22, y + .15, w - .4, .34, title_size, NAVY, True)
        textbox(slide, body, x + .22, y + .58, w - .4, h - .7, body_size, INK)
    return shape


def add_table(slide, x, y, widths, rows, row_h=.52, font_size=14,
              header_fill=NAVY, highlight_rows=()):
    table_shape = slide.shapes.add_table(len(rows), len(widths), Inches(x), Inches(y),
                                         Inches(sum(widths)), Inches(row_h * len(rows)))
    table = table_shape.table
    for col, width in enumerate(widths):
        table.columns[col].width = Inches(width)
    for r, row in enumerate(rows):
        for c, value in enumerate(row):
            cell = table.cell(r, c)
            cell.text = str(value)
            cell.margin_left = Inches(.09)
            cell.margin_right = Inches(.07)
            cell.margin_top = Inches(.035)
            cell.margin_bottom = Inches(.025)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            cell.fill.solid()
            if r == 0:
                cell.fill.fore_color.rgb = header_fill
            elif r in highlight_rows:
                cell.fill.fore_color.rgb = PALE
            else:
                cell.fill.fore_color.rgb = WHITE
            for paragraph in cell.text_frame.paragraphs:
                paragraph.font.name = "Aptos"
                paragraph.font.size = Pt(font_size)
                paragraph.font.bold = r == 0 or r in highlight_rows
                paragraph.font.color.rgb = WHITE if r == 0 else INK
    return table_shape


def add_image_fit(slide, path, x, y, w, h):
    with Image.open(path) as source:
        ratio = source.width / source.height
    box_ratio = w / h
    if ratio > box_ratio:
        width = w
        height = w / ratio
    else:
        height = h
        width = h * ratio
    return slide.shapes.add_picture(str(path), Inches(x + (w - width) / 2),
                                    Inches(y + (h - height) / 2),
                                    width=Inches(width), height=Inches(height))


def add_strip(slide, sheet_path, region, feature, x, y, w, h):
    # Keep the first held-out category row and epoch labels as a compact visual.
    with Image.open(sheet_path) as source:
        strip = source.crop((180, 0, min(source.width, 180 + 6 * 160), 210))
        buffer = io.BytesIO()
        strip.save(buffer, format="JPEG", quality=94)
    textbox(slide, f"{region} · feature {feature}", x, y, w, .26, 13, NAVY, True)
    buffer.seek(0)
    slide.shapes.add_picture(buffer, Inches(x), Inches(y + .28), width=Inches(w), height=Inches(h - .28))


# Cover
slide = prs.slides.add_slide(blank)
slide.background.fill.solid()
slide.background.fill.fore_color.rgb = NAVY
for x, y, w, h, color in [(.0, 0, .17, H, TEAL), (10.9, 0, 2.45, H, RGBColor(30, 53, 74))]:
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()
textbox(slide, "LOCALIZING FACE-RELATED\nACTIVATIONS IN A PRETRAINED SAE", .8, 1.25, 9.5, 1.6,
        31, WHITE, True)
textbox(slide, "Face-region fine-tuning and held-out feature localization", .84, 3.18, 9.0, .5,
        19, RGBColor(202, 220, 228))
textbox(slide, "Research update · October 2026", .84, 6.48, 8.6, .3, 13,
        RGBColor(202, 220, 228))
notes(slide, "This presentation summarizes two completed exploratory experiments: broad face-box fine-tuning and a separate eye-box fine-tuning run. It then presents region-specific feature localization on held-out images. The work measures SAE activation reconstruction and spatial association; it does not directly measure fairness or prove semantic concepts.")

# 2 — question and pipeline
slide = new_slide("Research question and experiment", "Can spatially focused fine-tuning improve reconstruction and reveal where SAE features activate?",
                  "Explain that we used a pretrained Top-K SAE on the SDXL U-Net activation block. The generator stays fixed. A copy of the SAE is fine-tuned using residual vectors inside detector boxes, then evaluated on held-out samples. The feature analyses select on validation data and evaluate on test data.")
steps = [
    ("SDXL portrait", "Generated image and saved U-Net activations", RGBColor(236, 241, 247)),
    ("Automatic boxes", "Florence-2 face and facial-part proposals", PALE),
    ("Pretrained SAE", "5,120 latent features · Top-K sparse codes", RGBColor(236, 241, 247)),
    ("Fine-tune / evaluate", "Face-box and eye-box runs · held-out test", PALE),
]
for i, (title, body, fill) in enumerate(steps):
    x = .8 + i * 3.08
    card(slide, x, 2.15, 2.72, 1.43, title, body, fill=fill, body_size=13)
    if i < len(steps) - 1:
        textbox(slide, "→", x + 2.76, 2.55, .3, .4, 23, TEAL, True, PP_ALIGN.CENTER)
card(slide, .8, 4.18, 5.7, 1.45, "What changed", "Only the SAE weights changed. No new latent nodes were added; the diffusion generator stayed frozen.", fill=WHITE)
card(slide, 6.75, 4.18, 5.7, 1.45, "What we measure", "Reconstruction MSE and feature activation inside annotated regions. These are not direct fairness or image-editing outcomes.", fill=WHITE, accent=ORANGE)

# 3 — data
slide = new_slide("Data and annotation setup", "The same 1,200-image portrait set supports the face and eye SAE experiments.",
                  "The source set contains six race prompts crossed with two gender prompts, 12 prompt categories total. It has 1,200 images: 720 train, 240 validation, 240 test. Florence-2 proposals are noisy and unreviewed. The 16-by-16 activation grid is much coarser than an image; box boundaries can include nearby content or miss small parts.")
card(slide, .85, 1.95, 3.65, 1.48, "1,200 portraits", "12 prompt categories\n720 train · 240 validation · 240 test", fill=WHITE)
card(slide, 4.83, 1.95, 3.65, 1.48, "Spatial supervision", "Automatic Florence-2 boxes\nMapped to a 16 × 16 activation grid", fill=WHITE, accent=ORANGE)
card(slide, 8.8, 1.95, 3.65, 1.48, "Held-out evaluation", "Feature selection on validation\nFinal measurements on test", fill=WHITE)
textbox(slide, "Important annotation limits", .9, 4.0, 4.0, .35, 18, NAVY, True)
textbox(slide, "• Boxes are automatic, not human-verified.\n• Small facial parts are detected inconsistently.\n• Missing boxes mean missing visual examples, not absence of the part.",
        .95, 4.48, 11.3, 1.45, 18, INK)

# 4 — broad face SAE results
slide = new_slide("Face-box fine-tuning improved held-out reconstruction", "Best checkpoint: epoch 20 · 240 held-out test images · lower MSE is better",
                  "The face run fine-tuned SAE activations at cells inside Florence-2 human-face boxes. On the test split, face-box MSE fell from 2.188 to 1.570, a 28.3% reduction. Outside-box MSE also improved from 2.133 to 1.769. The whole-grid score fell from 2.149 to 1.711. Because the SAE is reconstructing activations, not RGB images, these numbers do not by themselves say that generated faces look better.")
chart = ROOT / "sae-training/runs/face_sae_020/training_and_test_metrics.png"
add_image_fit(slide, chart, .65, 1.82, 8.15, 4.83)
card(slide, 9.05, 2.0, 3.55, 1.18, "Inside face box", "2.188 → 1.570\n28.3% lower", fill=WHITE, accent=TEAL, body_size=16)
card(slide, 9.05, 3.43, 3.55, 1.18, "Outside face box", "2.133 → 1.769\n17.1% lower", fill=WHITE, accent=ORANGE, body_size=16)
card(slide, 9.05, 4.86, 3.55, 1.18, "Whole activation grid", "2.149 → 1.711\n20.4% lower", fill=WHITE, body_size=16)
textbox(slide, "Reconstruction is measured in SAE activation space, not pixel space.", .9, 6.75, 11.5, .25, 12, MUTED)

# 5 — eye run
slide = new_slide("Eye-focused fine-tuning also improved test reconstruction", "Separate 20-epoch run · selected by validation eye MSE",
                  "This was a separate experiment initialized from the original pretrained SAE. Training loss used only activation vectors inside accepted eye boxes, with race/gender sampling balanced by category. The test comparison used 90 images with usable eye boxes. The eye MSE reduction was 16.2%. Across all 240 face-boxed test images, whole-face MSE improved about 10.1%. These remain activation reconstruction metrics; the run did not generate new RGB images.")
add_table(slide, .85, 2.0, [3.4, 2.55, 2.55, 2.0], [
    ["Test metric", "Original SAE", "Eye-fine-tuned", "Change"],
    ["Eye-box MSE · 90 images", "2.1007", "1.7608", "−16.2%"],
    ["Whole face · same 90", "2.1485", "1.9282", "−10.3%"],
    ["Whole face · all 240", "2.1881", "1.9671", "−10.1%"],
    ["All cells · all 240", "2.1491", "2.0307", "−5.5%"],
], row_h=.62, font_size=15, highlight_rows=(1,))
card(slide, .9, 5.48, 5.5, .95, "Training scope", "Updates SAE weights; does not train the image generator or add features.", fill=WHITE, body_size=13)
card(slide, 6.75, 5.48, 5.5, .95, "Interpretation", "Better activation reconstruction does not prove eye-specific semantics.", fill=WHITE, accent=ORANGE, body_size=13)

# 6 — across epoch example
slide = new_slide("Activation maps across training checkpoints", "Same held-out portraits · same feature · epochs 0, 4, 8, 12, 16, and 20",
                  "This sheet follows feature 3701 across checkpoints. It is selected as an eye-associated feature, but this broad face-box view shows where it activates on full face regions. The overlay is normalized per image; compare the location of activity more than raw color strength between different images.")
face_map = ROOT / "sae-training/runs/eye_sae_001/region_heatmaps/face/feature_3701.jpg"
add_image_fit(slide, face_map, .85, 1.78, 11.7, 4.95)
textbox(slide, "Cyan = automatic face box · `jet` colors = feature activation inside the box", .9, 6.82, 11.3, .25, 12, MUTED)

# 7 — region-specific montage
slide = new_slide("Each part now has its own feature candidates", "Examples across epochs, with the activation overlay restricted to that part’s annotation box",
                  "These compact strips use one Black Man held-out example per region to show the same sample across all six checkpoints. Full category-by-epoch sheets are saved in the run folder. Different parts use different features selected on validation data. Nose/ear boxes are especially sparse; this slide is an illustration, not a population-level comparison.")
specific = ROOT / "sae-training/runs/eye_sae_001/region_specific_heatmaps"
panels = [("eyes", 3701), ("eyebrows", 1600), ("nose", 1664), ("ears", 4858)]
positions = [(.8, 1.78), (6.9, 1.78), (.8, 4.25), (6.9, 4.25)]
for (region, feature), (x, y) in zip(panels, positions):
    add_strip(slide, specific / region / f"feature_{feature}.jpg", region.title(), feature,
              x, y, 5.65, 1.75)
textbox(slide, "Full sheets: 9/12 categories for eyes, 7/12 eyebrows, 4/12 nose, 6/12 ears.",
        .85, 6.55, 11.6, .3, 12, MUTED)

# 8 — quantitative localization
slide = new_slide("Selected features activate more inside their part boxes", "Validation selected candidates; test set measured mean part-minus-reference activation",
                  "The contrast is mean activation in the part box minus mean activation in a reference area. It is not a correlation coefficient. Candidate ranking is done on validation data; confidence intervals are stratified image-bootstrap intervals on the held-out test split. Positive intervals support spatial association with these automatic boxes, not semantic identity. Ear comparison uses the rest of the image because ears can lie outside the face box.")
add_table(slide, .62, 1.82, [1.55, 2.55, 1.35, 1.5, 4.4], [
    ["Region", "Top features", "Test n", "Categories", "Test contrast range (95% CI range)"],
    ["Eyes", "3701, 3593, 1600", "90", "9", "+6.80 to +9.03\n[+6.30, +10.66]"],
    ["Eyebrows", "1600, 3701, 1237", "28", "7", "+6.27 to +7.96\n[+4.27, +9.70]"],
    ["Nose", "1664, 3593, 1991", "13", "4", "+6.54 to +11.79\n[+5.08, +12.26]"],
    ["Ears", "4858, 2230, 2611", "18", "6", "+5.39 to +10.18\n[+2.43, +11.40]"],
], row_h=.78, font_size=12, highlight_rows=(1,))
card(slide, .85, 6.12, 5.65, .7, "Meaning", "Evidence of localization to annotated boxes.", fill=PALE, body_size=12)
card(slide, 6.78, 6.12, 5.65, .7, "Not established", "Dedicated semantics, causal control, or fairness effects.", fill=WHITE, accent=ORANGE, body_size=12)

# 9 — broader feature-to-part associations
slide = new_slide("Other intervention features also associate with face parts",
                  "19 fixed feature IDs · original pretrained SAE · held-out test portraits",
                  "The feature IDs came from the existing feature-intervention runs and were fixed before this held-out analysis. Each score is mean activation inside the automatic part box minus activation in the reference area, with a stratified image-bootstrap 95% interval. The same feature can appear in multiple rows because several features associate with more than one part. Feature 1664 has the clearest nose-centered pattern, while 527 and 3725 associate with several regions. These are spatial associations, not proof of semantic identity or causal editing. The nose and ear estimates have sparse sample coverage, and intervals are unadjusted across 76 feature-region comparisons.")
add_table(slide, .72, 1.78, [1.45, 7.6, 2.05], [
    ["Region", "Leading fixed features · contrast [95% interval]", "Test images / categories"],
    ["Eyes", "527: +1.75 [1.40, 2.11] · 1664: +1.65 [0.85, 2.42] · 3725: +0.78 [0.43, 1.23]", "90 / 9"],
    ["Eyebrows", "527: +2.66 [2.25, 3.04] · 3725: +1.62 [1.33, 1.91] · 3722: +0.61 [0.50, 0.71]", "28 / 7"],
    ["Nose", "1664: +12.04 [11.48, 12.56] · 3725: +0.86 [0.70, 1.03] · 527: +0.43 [0.19, 0.70]", "13 / 4"],
    ["Ears", "527: +5.60 [4.70, 6.49] · 3725: +4.10 [3.12, 5.02] · 3722: +2.89 [1.65, 4.09]", "18 / 6"],
], row_h=.84, font_size=11, highlight_rows=(3,))
card(slide, .85, 6.2, 5.65, .68, "Clearest part-centered pattern", "Feature 1664: strong nose contrast; lower or negative contrast elsewhere.", fill=PALE, body_size=11)
card(slide, 6.78, 6.2, 5.65, .68, "Broader associations", "Features 527 and 3725 show positive contrast in several parts.", fill=WHITE, accent=ORANGE, body_size=11)
textbox(slide, "Contrasts are SAE activation units, not percentages or correlation coefficients.",
        .88, 6.96, 11.4, .2, 10, MUTED)

# 10 — qualitative heatmap examples
slide = new_slide("Jet heatmaps show example feature locations",
                  "Held-out portraits · cyan = detected face · yellow = annotated part",
                  "These are illustrative high-activation examples, selected from eligible held-out test portraits for display. The heatmap is normalized within each detected face box. The examples help inspect spatial alignment, but do not estimate how often a feature localizes to a part; those estimates are in the preceding table and use all eligible test images. Boxes are automatic and unreviewed.")
examples = ROOT / "sae-training/runs/eye_sae_001/part_feature_validation_intervention19_baseline/images/feature_part_examples.jpg"
add_image_fit(slide, examples, .65, 1.68, 5.65, 5.25)
textbox(slide, "Feature 1664", 6.65, 1.9, 5.7, .38, 19, NAVY, True)
textbox(slide, "The strongest activation in this example falls around the annotated nose box, consistent with its high held-out nose contrast.",
        6.67, 2.38, 5.65, 1.05, 16, INK)
textbox(slide, "Features 527 and 3725", 6.65, 3.85, 5.7, .38, 19, NAVY, True)
textbox(slide, "Their heatmaps spread across more of the face, matching their positive contrasts in several regions rather than a single-part pattern.",
        6.67, 4.33, 5.65, 1.05, 16, INK)
card(slide, 6.65, 5.85, 5.55, .82, "How to read these panels", "Examples selected for visibility; aggregate table uses the full held-out sample.", fill=PALE, body_size=12)

# 11 — caveats / next steps
slide = new_slide("What we can claim—and what comes next", "The evidence supports a useful localization result, with clear limits",
                  "Close with a careful summary. Both fine-tuning runs improved held-out SAE activation reconstruction. The region-specific candidates show positive held-out activation contrast. However, the automatic boxes are unreviewed and coverage varies sharply; nose has only 13 test images in four categories. The next steps are human review of part annotations, more usable nose/ear boxes, then evaluating interventions with direct outcome measures before making fairness claims.")
card(slide, .85, 1.95, 5.6, 1.5, "Supported by these experiments", "• Face and eye fine-tuning lowered held-out activation MSE.\n• Candidate features had positive held-out part-box contrast.", fill=PALE, body_size=14)
card(slide, 6.75, 1.95, 5.6, 1.5, "Still unknown", "• Whether a feature has a stable semantic meaning.\n• Whether changing it edits only that part.\n• Whether any fairness outcome improves.", fill=WHITE, accent=ORANGE, body_size=14)
textbox(slide, "Next steps", .9, 4.05, 2.0, .36, 18, NAVY, True)
textbox(slide, "1. Audit detector boxes, especially nose and ear annotations.\n2. Expand part coverage and repeat held-out localization.\n3. If testing interventions, measure target presence and collateral changes directly.\n4. Evaluate any fairness claim with a separate, explicit fairness outcome.",
        .95, 4.55, 11.4, 1.7, 17, INK)

# 12 — headwear intervention result
slide = new_slide("Negative steering did not remove headwear cleanly",
                  "Same 20 Middle Eastern woman prompt images · seeds 0–19 · α +1 = unchanged",
                  "The feature 138 and 3725 interventions use the same SDXL Turbo U-Net block and SAE checkpoint. We visually compared same-seed originals with alpha -3 and alpha -50 outputs. For feature 138, mild negative steering often adds or exaggerates ornate headwear; strong steering produces abstract metallic texture. For feature 3725, mild negatives leave scarves in place; strong steering creates large beige fabric-like shapes and removes or fragments faces. This is a qualitative review by one reviewer, not a blinded study. Florence-2 pass counts do not measure face preservation or clean attribute removal.")

feature_runs = [
    ("Feature 138", "Mild negatives add ornate headwear; stronger values destroy the face.",
     ROOT / "sae-activation-tuning/tuning_activations/feature_138/run_biased_002/images/middle_eastern_woman/baseline/seed000.png",
     ROOT / "sae-activation-tuning/tuning_activations/feature_138/run_biased_003_mild_negative/images/middle_eastern_woman/seed000_alpha_m3.png",
     ROOT / "sae-activation-tuning/tuning_activations/feature_138/run_biased_002/images/middle_eastern_woman/seed000_alpha_m50.png"),
    ("Feature 3725", "Mild negatives leave scarves; stronger values obscure the face.",
     ROOT / "sae-activation-tuning/tuning_activations/feature_3725/run_biased_slurm_4109210_1/images/middle_eastern_woman/baseline/seed000.png",
     ROOT / "sae-activation-tuning/tuning_activations/feature_3725/run_biased_slurm_4109210_1/images/middle_eastern_woman/seed000_alpha_m3.png",
     ROOT / "sae-activation-tuning/tuning_activations/feature_3725/run_biased_slurm_4109210_1/images/middle_eastern_woman/seed000_alpha_m50.png"),
]
for idx, (name, summary, original, mild, strong) in enumerate(feature_runs):
    x = .7 + idx * 6.05
    panel = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(1.72), Inches(5.9), Inches(3.67))
    panel.fill.solid(); panel.fill.fore_color.rgb = WHITE
    panel.line.color.rgb = RGBColor(222, 230, 236)
    accent_bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(1.72), Inches(.07), Inches(3.67))
    accent_bar.fill.solid(); accent_bar.fill.fore_color.rgb = TEAL if idx == 0 else ORANGE
    accent_bar.line.fill.background()
    textbox(slide, name, x + .22, 1.86, 5.5, .34, 17, NAVY, True)
    textbox(slide, summary, x + .22, 2.26, 5.5, .36, 11, INK)
    labels = [("Original", original), ("α = −3", mild), ("α = −50", strong)]
    for j, (label, path) in enumerate(labels):
        tile_x = x + .18 + j * 1.88
        textbox(slide, label, tile_x, 2.72, 1.72, .22, 10, MUTED, True,
                align=PP_ALIGN.CENTER)
        add_image_fit(slide, path, tile_x, 2.98, 1.72, 1.72)

card(slide, .7, 5.65, 11.95, 1.0, "Conclusion",
     "Neither feature produced clean headwear-only removal in this pilot. Count a result as successful only if the face stays intact and the headwear changes without major artifacts.",
     fill=PALE, accent=TEAL, title_size=15, body_size=13)
textbox(slide, "Qualitative review of 20 generated images in one prompt group; one reviewer. Feature IDs are model-specific candidates, not verified semantic labels.",
        .78, 6.78, 11.5, .22, 9, MUTED)

OUT.parent.mkdir(parents=True, exist_ok=True)
prs.save(OUT)
print(OUT)
