"""Build an editable research presentation and layout previews from local findings."""
import sys
sys.path.insert(0, '/tmp/sae_ppt_dependencies')
from pathlib import Path
import csv
import json
import math
from collections import defaultdict
from zipfile import ZipFile
import xml.etree.ElementTree as ET

from PIL import Image, ImageDraw, ImageFont, ImageOps
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE

ROOT = Path('/n/fs/goose')
DATA = ROOT / 'baseline/sae_bias_interventions_20261009'
OUT = ROOT / 'baseline/sae_bias_presentation'
OUT.mkdir(exist_ok=True)
(OUT / 'previews').mkdir(exist_ok=True)
BG, INK, MUTED, TEAL, ORANGE, PALE, LINE = '#F7F8FA', '#152B42', '#59697A', '#087E8B', '#D97936', '#E8F2F3', '#DCE3E9'
SCALE = 120
W, H = 13.333, 7.5
prs = Presentation()
prs.slide_width, prs.slide_height = Inches(W), Inches(H)
prs.core_properties.title = 'Sparse Autoencoders for Auditing Racial Bias in SDXL Turbo'
prs.core_properties.subject = 'Controlled feature interventions on existing portrait baselines'
prs.core_properties.author = 'Research experiment'
prs.core_properties.keywords = 'SDXL Turbo, sparse autoencoder, racial bias, interpretability, feature intervention'
previews, notes = [], []


def rgb(color):
    return RGBColor.from_string(color.lstrip('#'))


def font(size, bold=False):
    filename = 'DejaVuSans-Bold.ttf' if bold else 'DejaVuSans.ttf'
    return ImageFont.truetype('/usr/share/fonts/dejavu-sans/' + filename, round(size * SCALE / 72))


def box(slide, x, y, w, h, color=PALE, border=None):
    sh = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    sh.fill.solid()
    sh.fill.fore_color.rgb = rgb(color)
    if border:
        sh.line.color.rgb = rgb(border)
    else:
        sh.line.fill.background()
    canvas = slide._preview
    ImageDraw.Draw(canvas).rectangle((x*SCALE,y*SCALE,(x+w)*SCALE,(y+h)*SCALE), fill=color, outline=border)
    return sh


def wrap(text, f, maxwidth):
    result = []
    for paragraph in text.split('\n'):
        words = paragraph.split()
        line = ''
        for word in words:
            trial = (line + ' ' + word).strip()
            if f.getlength(trial) > maxwidth and line:
                result.append(line)
                line = word
            else:
                line = trial
        result.append(line)
    return result


def txt(slide, text, x, y, w, h, size=20, color=INK, bold=False):
    # Fit using the same conservative font metrics as the layout preview.
    while True:
        f = font(size, bold)
        paragraphs = text.split('\n')
        line_count = sum(len(wrap(line, f, (w-.04)*SCALE)) for line in paragraphs)
        estimate = 2 + line_count*size*SCALE/72*1.22 + max(0,len(paragraphs)-1)*8*SCALE/72
        if estimate <= h*SCALE+8:
            break
        size -= .5
        if size < 10:
            raise ValueError(f'Cannot fit text: {text[:65]}')
    sh = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = sh.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(.02)
    tf.margin_top = tf.margin_bottom = Inches(.01)
    for i, line in enumerate(text.split('\n')):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = line
        p.font.name = 'DejaVu Sans'
        p.font.size = Pt(size)
        p.font.bold = bold
        p.font.color.rgb = rgb(color)
        p.space_after = Pt(8)
    draw = ImageDraw.Draw(slide._preview)
    f = font(size, bold)
    yy = y*SCALE + 2
    for paragraph in text.split('\n'):
        lines = wrap(paragraph, f, (w-.04)*SCALE)
        for line in lines:
            draw.text((x*SCALE+2, yy), line, font=f, fill=color)
            yy += size*SCALE/72*1.22
        yy += 8*SCALE/72
    # Conservative check against overflowing text in generated previews.
    if yy - 8*SCALE/72 > (y+h)*SCALE + 8:
        raise ValueError(f'Text exceeds box on slide {len(prs.slides)}: {text[:65]}')
    return sh


def pic(slide, path, x, y, w, h):
    path = Path(path)
    with Image.open(path) as im:
        iw, ih = im.size
        ratio = min(w/iw, h/ih)
        pw, ph = iw*ratio, ih*ratio
        xx, yy = x+(w-pw)/2, y+(h-ph)/2
        slide.shapes.add_picture(str(path), Inches(xx), Inches(yy), width=Inches(pw), height=Inches(ph))
        preview = im.convert('RGB').resize((round(pw*SCALE), round(ph*SCALE)), Image.Resampling.LANCZOS)
        slide._preview.paste(preview, (round(xx*SCALE), round(yy*SCALE)))


def slide(title, subtitle, section, note, source='FINDINGS.md; sae_bias_experiment.py'):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    s.background.fill.solid()
    s.background.fill.fore_color.rgb = rgb(BG)
    s._preview = Image.new('RGB', (round(W*SCALE),round(H*SCALE)), BG)
    box(s, .45, .4, .08, .48, TEAL)
    txt(s, section.upper(), .7, .32, 11.8, .3, size=10, color=TEAL, bold=True)
    txt(s, title, .7, .84, 12, .78, size=29, bold=True)
    if subtitle:
        txt(s, subtitle, .72, 1.65, 11.9, .65, size=15, color=MUTED)
    box(s, .7, 7.02, 11.9, .012, LINE)
    txt(s, 'SDXL Turbo · SAE portrait audit | exploratory findings', .72, 7.1, 7, .23, size=9, color=MUTED)
    txt(s, f'{len(prs.slides):02d}', 12.12, 7.08, .42, .28, size=11, color=TEAL, bold=True)
    s.notes_slide.notes_text_frame.text = note + '\n\nLocal evidence: ' + source
    notes.append((title,note,source))
    previews.append(s._preview)
    return s


def card(s, x, y, w, h, title, body, accent=TEAL):
    box(s,x,y,w,h,'#FFFFFF',LINE)
    box(s,x,y,.045,h,accent)
    txt(s,title,x+.2,y+.16,w-.4,.65,size=19,bold=True,color=accent)
    txt(s,body,x+.2,y+.96,w-.4,h-1.05,size=17)


def label_image(s, path, label, x,y,w=2.65,h=2.65):
    txt(s,label,x,y,w,.32,size=12,color=TEAL,bold=True)
    pic(s,path,x,y+.4,w,h)


s=slide('Sparse autoencoders for a racial-bias audit','Controlled interventions on an existing SDXL Turbo portrait baseline','Research briefing',
    'This presentation describes work that was actually performed on local artifacts. The central finding is methodological: the pipeline supports exact matched-seed replay and controlled feature scaling. It does not establish that racial bias or image quality improved. The portrait study uses the existing pretrained SAE so old feature IDs retain their reference. A separate small SAE training pilot is included only in the appendix. All race and gender labels refer to generation prompts, not inferred identities.')
txt(s,'Interpretable controls are working.\nBias reduction is still an open question.',.8,2.65,11.7,1.7,size=30,bold=True)
for x,value,desc in [(0.8,'600','existing baseline portraits'),(4.85,'432','controlled feature edits'),(8.9,'0','pixel error in replay / no-op checks')]:
    box(s,x,4.85,3.65,1.35,PALE)
    txt(s,value,x+.18,5.02,3.2,.58,size=30,color=TEAL,bold=True)
    txt(s,desc,x+.18,5.72,3.25,.38,size=12,color=MUTED)

s=slide('Three questions require different evidence','A feature can distinguish prompt groups without controlling an unwanted attribute.','Research question',
    'We distinguish association, intervention, and fairness evaluation. Association measures whether SAE activations vary with generation-prompt categories. The intervention tests what changes when a specific internal feature contribution is multiplied while holding the prompt and random seed fixed. A fairness claim additionally requires a predefined unwanted outcome, annotations, and a comparison of outcome rates across groups. The current run completes the first two stages for selected features; it has not completed the third.')
card(s,.8,2.6,3.85,3.45,'1 · Association','Do feature activations differ across the six prompt groups?\nMeasured on cached SAE vectors.')
card(s,4.75,2.6,3.85,3.45,'2 · Intervention','What changes when one feature is suppressed or amplified?\nTested with fixed seeds.')
card(s,8.7,2.6,3.85,3.45,'3 · Bias outcome','Does an unwanted outcome decrease without collateral damage?\nNot yet measured.',ORANGE)

s=slide('Two existing datasets, two distinct roles','The cached SAE vectors and the saved portrait PNGs are separate artifacts.','Data',
    'The baseline has 50 images per race/gender category, for 600 PNGs total. The separate cached feature file also contains 600 samples, each represented by 5,120 spatial-max-pooled pretrained SAE activations. It uses randomly sampled seeds from its own generation script. That file is used for feature discovery and held-out association reporting; baseline seeds 40, 41, and 42 are used for visual interventions. The cached codes are not the raw 1,280-channel residual inputs needed to train a new SAE. The cached file does not store per-row seed metadata.',
    'baseline/images/; race_classifier/race_classify_features.npz; race_classifier/generate_data.py')
card(s,.8,2.55,5.72,3.45,'Baseline images','6 prompt races × 2 genders × 50 seeds\n600 saved portraits\nReplay sample: seeds 40, 41, 42\n36 matched baseline images')
card(s,6.72,2.55,5.72,3.45,'Cached feature vectors','600 images × 5,120 SAE dimensions\nSpatial maximum over each feature map\n480 discovery / 120 held-out samples\n40 / 10 samples in each category')
txt(s,'Labels are taken from generation prompts. No racial identity is inferred from image pixels.',.9,6.35,11.5,.4,size=16,color=TEAL,bold=True)

s=slide('The baseline spans all 12 prompt categories','Seed 40 shown in each cell; identical prompt template except for race and gender.','Baseline coverage',
    'This slide provides the category coverage, not an assessment of the depicted people’s identities. Rows correspond to man and woman prompts. Columns follow the six requested race labels. The template fixes expression, shirt color, background, and a request for no facial coverings. The original phrase does not explicitly exclude a headscarf that leaves the face visible. This distinction matters when interpreting the Middle Eastern woman example. We replay this exact template rather than changing it during the intervention experiment.',
    'baseline/images/*/baseline40.png; baseline/baseline-image-generation.ipynb')
races=['black','white','east asian','middle eastern','indian','hispanic']
for j,race in enumerate(races):
    x=.82+j*2.07
    txt(s,race.title(),x,2.35,1.95,.4,size=12,bold=True,color=TEAL)
    for i,gender in enumerate(['man','woman']):
        y=2.86+i*1.93
        pic(s,ROOT/'baseline/images'/f'{race.replace(" ","_")}_{gender}'/'baseline40.png',x,y,1.76,1.76)
        txt(s,gender,x,y+1.78,1.8,.24,size=10,color=MUTED)

s=slide('What the SAE represents','A sparse dictionary of intermediate U-Net residual changes—not input noise or image latents.','Model',
    'The analysis taps unet.down_blocks.2.attentions.1. For every spatial position, the training representation is the attention block output minus its input, with 1,280 channels. The existing SAE expands this into a 5,120-dimensional feature dictionary using Top-K encoding with k=10 and ReLU. Each decoder column is a learned residual direction. The pretrained checkpoint also has an auxiliary-k setting of 256 in its configuration. We do not retrain this checkpoint in the portrait audit. This preserves feature IDs from the previous notebooks.',
    'sdxl-unbox/SAE/sae.py; pretrained down.2.1 final/config.json; sae_bias_experiment.py')
for x,title,body in [(0.85,'Block residual','r = output − input\n1,280 channels per position'),(5.0,'Sparse encoding','z = ReLU(TopK(Wₑ(r − b) + c))\n5,120 features; k = 10'),(9.05,'Reconstruction','r̂ = Wᵈz + b\nLearned decoder directions')]:
    card(s,x,2.8,3.45,2.65,title,body)
txt(s,'→',4.42,3.7,.5,.8,size=32,color=TEAL)
txt(s,'→',8.57,3.7,.5,.8,size=32,color=TEAL)
txt(s,'The same pretrained dictionary is used for every group; feature IDs are checkpoint-specific.',.9,5.95,11.6,.6,size=18,color=MUTED)

s=slide('Feature scaling preserves reconstruction error','Only the selected feature contribution is changed; the remaining block output stays intact.','Intervention equation',
    'Let r be the original residual and r-hat its SAE reconstruction. The residual reconstruction error is e = r minus r-hat. Scaling feature f from z-f to s times z-f changes the reconstruction by (s minus 1) times z-f times decoder direction d-f. Adding this delta directly to the original block output preserves e. This prevents generic SAE reconstruction artifacts from being mixed into every intervention. Strength one is therefore a mathematical no-op. Inactive features remain inactive under multiplicative editing; this procedure is not an additive feature injection.',
    'sdxl-unbox/scripts/sae_experiment.py: intervention; sae_bias_experiment.py')
box(s,.9,2.65,11.55,1.35,PALE)
txt(s,'output′ = output + (s − 1) · z_f · d_f',1.2,2.95,10.9,.73,size=30,color=TEAL,bold=True)
for x,strength,label in [(0.9,'s = 0','remove contribution'),(3.85,'s = 0.5','halve contribution'),(6.8,'s = 1','exact no-op'),(9.75,'s = 2','double contribution')]:
    txt(s,strength,x,4.65,2.7,.6,size=25,bold=True)
    txt(s,label,x,5.35,2.7,.58,size=15,color=MUTED)
txt(s,'Spatial activations are scaled separately at every location. Features are edited one at a time.',.9,6.25,11.5,.43,size=16,color=MUTED)

s=slide('Candidate ranking controls for prompt gender','Discovery and held-out association are computed separately.','Feature discovery',
    'For each race/gender cell, an RNG with seed zero produces a 40/10 discovery/held-out split. For a given race and feature, we calculate the difference between that race’s discovery mean and the mean of other races within each gender. We divide by the feature’s standard deviation across all discovery samples, clipped at 1e-6 for numerical stability, and average the two gender contrasts equally. Candidate 3530 is the strongest additional gender-sign-consistent discovery association. Features 138 and 527 were included because of prior analyses. Held-out effects do not select the candidate. This is exploratory; no multiple-testing-adjusted significance is reported.',
    'sae_bias_experiment.py: analyze; feature_associations.csv; cached_data_split.npz')
box(s,.85,2.65,11.6,1.25,PALE)
txt(s,'C(r,f) = mean_g [(μ(r,g,f) − μ(other races,g,f)) / σ(discovery,f)]',1.05,2.95,11.1,.65,size=22,color=TEAL,bold=True)
card(s,.85,4.22,5.65,2.2,'Discovery: 480 samples','Rank absolute standardized contrast.\nCheck sign agreement across genders.')
card(s,6.72,4.22,5.65,2.2,'Held-out: 120 samples','Report the same selected feature’s effect.\nNo confidence intervals or p-values.')

associations=list(csv.DictReader((DATA/'feature_associations.csv').open()))
top=[next(row for row in associations if row['prompt_race']==r) for r in races]
s=slide('Prompt-associated features persist on held-out data','Different features emerge for different prompt groups; these are associations, not semantic labels.','Association results',
    'The chart shows the strongest discovery association for each requested race label and its corresponding held-out contrast. The numerical values are standardized effect contrasts, not classification accuracy or bias scores. All six displayed features have a positive association with their own prompt label in both splits. Only 527 and 3530 from this list were intervened on; feature 138 is an additional previously studied candidate. The other four top features remain association findings. The same prompt template appears in discovery and held-out data, so this does not demonstrate transfer to new prompt formulations.',
    'feature_associations.csv')
for i,(race,row) in enumerate(zip(races,top)):
    yy=2.7+i*.55
    txt(s,f'{race.title()} · f{row["feature"]}',.9,yy,3.0,.36,size=14)
    a,b=float(row['train_standardized_contrast']),float(row['held_out_standardized_contrast'])
    box(s,4.05,yy+.01,a*2.55,.12,TEAL)
    box(s,4.05,yy+.18,b*2.55,.12,ORANGE)
    txt(s,f'{a:.2f} / {b:.2f}',11.15,yy,1.2,.4,size=13,color=MUTED)
txt(s,'Discovery',4.1,6.25,1.5,.35,size=13,color=TEAL,bold=True)
txt(s,'Held-out',5.65,6.25,1.5,.35,size=13,color=ORANGE,bold=True)
txt(s,'Standardized contrast; larger is stronger association',7.0,6.25,5.2,.35,size=12,color=MUTED)

s=slide('A fully matched intervention experiment','Three seeds, all categories, the same three features, and four strengths.','Experimental controls',
    'The experiment includes 12 categories, three seeds, three features, and four strengths, yielding 432 edited images. The three features are 138, 527, and 3530. Replay uses the original fp16 VAE, Euler ancestral scheduler with trailing timesteps, one denoising step, guidance scale zero, and separately seeded latent and scheduler random streams. The original baseline and the new replay matched exactly for all 36 cases. Every strength-one image was also pixel-identical to its replay. This validates reproducibility and the no-op control, but does not measure fairness or quality. The current run generated 36 category/feature grids.',
    'settings.json; intervention_metrics.csv; sae_bias_experiment.py')
txt(s,'12 categories × 3 seeds × 3 features × 4 strengths = 432 edits',.9,2.65,11.7,.75,size=25,bold=True,color=TEAL)
card(s,.9,3.75,5.65,2.25,'36 / 36 baseline replays','Pixel-identical to saved baseline PNGs.\nSeed values: 40, 41, 42.')
card(s,6.75,3.75,5.65,2.25,'108 / 108 no-op controls','Strength 1 is pixel-identical to replay.\nNo reconstruction replacement in edits.')
txt(s,'Checkpoint: down.2.1 · 5,120 features · k=10 | SDXL Turbo · 512² · 1 step · CFG=0',.95,6.4,11.5,.35,size=13,color=MUTED)


def visual_slide(title,subtitle,group,feature,seeds,note):
    s=slide(title,subtitle,'Visual intervention',note,
            f'{group}/feature{feature}_comparison.png; individual seed PNGs; intervention_metrics.csv')
    for row,seed in enumerate(seeds):
        y=2.38+row*2.12
        txt(s,f'Seed {seed}',.8,y+.84,.72,.45,size=12,bold=True,color=MUTED)
        choices=[('Original',f'seed{seed}_original.png'),('Activation map',f'seed{seed}_feature{feature}_heatmap.png'),
                 ('Strength 0',f'seed{seed}_feature{feature}_strength0.png'),('Strength 2',f'seed{seed}_feature{feature}_strength2.png')]
        for j,(label,file) in enumerate(choices):
            x=1.62+j*2.74
            txt(s,label,x,y,2.4,.3,size=12,bold=True,color=TEAL)
            pic(s,DATA/group/file,x,y+.38,2.38,1.68)
    return s

s=visual_slide('Feature 138 does not cleanly control head coverings','Middle Eastern woman prompts; scarf remains in seed 40 at strength 0.','middle_eastern_woman',138,[40,41],
    'The side-by-side images show two seeds from the Middle Eastern woman prompt category. Seed 40 contains a visible headscarf; removing feature 138’s contribution does not remove it. Seed 41 has visible uncovered hair, and the feature is active there as well. Scaling produces changes in face and hair detail, which agrees with the earlier notebook’s plain-hair confound. Heatmap intensity is normalized independently per image, so colors do not support cross-image magnitude comparisons. The original prompt prohibits facial coverings, not all head coverings; this example is not by itself an established prompt violation.')

s=visual_slide('Feature 3530 is associated, but not a scarf switch','The strongest additional candidate changes detail without removing the seed-40 scarf.','middle_eastern_woman',3530,[40,42],
    'Feature 3530 was selected as the strongest additional gender-sign-consistent discovery association after including 138 and 527. Its held-out standardized contrast for the Middle Eastern prompt label is 2.65. Despite that association, the intervention at strength zero leaves the scarf in seed 40. The visible changes are smaller and involve facial or local detail in these examples. This separates prompt association from control of a specific attribute. It does not imply the feature lacks effects under other prompts, strengths, blocks, or denoising steps; those were not tested here.')

s=visual_slide('Feature 527 changes depiction, not an evaluated bias score','Black man prompts; changes include texture/contrast and pose or crop at some settings.','black_man',527,[40,42],
    'The two examples use the Black man prompt category. Feature 527 has a strong discovery and held-out association with the Black prompt label. Reducing or amplifying its contribution changes face texture and contrast. In seed 42, amplification also produces visible differences in composition and facial detail. These visual observations do not establish a semantic feature label or a quality preference. Suppressing a group-associated feature can alter a requested depiction without improving fairness. No identity inference or race classifier is applied to the edited images in this experiment.')

s=slide('Intervention magnitude varies across prompt groups','Mean absolute RGB pixel change at strength 0; averaged over both genders and three seeds.','Measured image effects',
    'The table aggregates the recorded mean absolute pixel differences between original and edited PNGs at strength zero. Each cell averages six images, two genders times three seeds. Values are on an 8-bit RGB scale from zero to 255. This quantifies whether an intervention has an effect and how large its image-level change is, not whether it improves quality or fairness. Feature 138 and 3530 are inactive in several tested categories, giving zero changes. Feature 527 also changes some Hispanic, Indian, and Middle Eastern samples. Results are descriptive with a small sample and no uncertainty intervals.',
    'intervention_metrics.csv; test_activation_metrics.csv')
metrics=list(csv.DictReader((DATA/'intervention_metrics.csv').open()))
buckets=defaultdict(list)
for row in metrics:
    if float(row['strength'])==0:
        buckets[(row['prompt_race'],row['feature'])].append(float(row['pixel_change']))
txt(s,'Prompt group',.95,2.55,3.5,.45,size=17,bold=True)
for j,f in enumerate([138,527,3530]): txt(s,f'Feature {f}',4.45+j*2.65,2.55,2.3,.45,size=17,bold=True,color=TEAL)
for i,race in enumerate(races):
    y=3.15+i*.47
    txt(s,race.title(),.95,y,3.2,.38,size=16)
    for j,f in enumerate([138,527,3530]):
        values=buckets[(race,str(f))]
        value=sum(values)/len(values)
        x=4.45+j*2.65
        box(s,x-.08,y-.02,2.2,.41,PALE if value>0 else '#EEF0F3')
        txt(s,f'{value:.2f}',x,y,2.0,.38,size=16,color=TEAL if value>0 else MUTED,bold=value>0)
txt(s,'0 means no image change under this intervention—not absence of bias. Pixel change is not a quality metric.',.95,6.35,11.5,.5,size=15,color=MUTED)

s=slide('What the evidence supports—and what it does not','The main contribution is a reproducible feature-intervention audit.','Interpretation',
    'The experiments support exact baseline replay, correct multiplicative editing, associations that recur in a held-out cached split, and qualitative changes from selected feature interventions. They do not establish semantic race features, headscarf-specific control, mitigation of racial bias, improvement in image quality, or generalization beyond the current prompt template. An older classifier achieved perfect held-out classification on its own split, but this is prompt-group decodability, not a fairness score, and no classifier was applied to edited images in the present run. The distinction should be central to the research narrative.')
card(s,.9,2.6,5.65,3.65,'Supported','Exact saved-baseline replay\nCorrect no-op and strength scaling\nHeld-out prompt-group associations\nFeature-dependent image changes')
card(s,6.75,2.6,5.65,3.65,'Not established','A feature’s semantic meaning\nHead-covering removal\nReduced racial bias\nImproved quality or prompt transfer',ORANGE)

s=slide('Threats to validity','A useful exploratory result requires a narrower claim than “we removed bias.”','Limitations',
    'The intervention sample has only three seeds per category and uses one template, one block, and one denoising step. Candidate selection is partly informed by older analyses, so the held-out cached split is exploratory rather than a pristine confirmatory test. The spatial-max feature representation can emphasize a single location and does not directly locate a semantic concept. The intervention is restricted to features that already activate. Human reviews remain blank, and no inter-rater agreement, confidence intervals, or downstream quality measures have been computed. Overlay maps are normalized independently and should not be used for magnitude comparisons.')
card(s,.85,2.6,3.85,3.65,'Coverage','3 seeds per category\n1 prompt template\n1 U-Net block\n1 denoising step')
card(s,4.75,2.6,3.85,3.65,'Measurement','Spatial-max associations\nNo outcome annotations\nNo quality scores or intervals\nPer-image heatmap scaling')
card(s,8.65,2.6,3.85,3.65,'Selection','Exploratory candidate choices\nPrior notebook knowledge\nNo held-out prompt family\nInactive features stay inactive',ORANGE)

s=slide('The next study should target a defined unwanted outcome','Use paired prompts and hold out seeds before feature selection.','Proposed next experiment',
    'This slide proposes future work, not completed results. For a head-covering target, use a common explicit no-head-covering instruction for every category, along with a paired prompt variant. Reserve common seeds across all categories for final evaluation before candidate selection. Collect residual vectors from multiple blocks and compare matched activation changes, then evaluate feature interventions on those untouched seeds. Annotate unwanted coverings, hair damage, clothing compliance, image artifacts, and image preference with blinded or randomized presentation. Report within-group outcome changes and between-group disparities with uncertainty. Prompt compliance and requested appearance should be assessed alongside the bias target.')
for x,n,title,body in [(0.85,'01','Define the outcome','Explicit prompt constraint;\nconsistent across groups.'),(4.85,'02','Select specific controls','Matched prompts / seeds;\ncompare multiple blocks.'),(8.85,'03','Evaluate on untouched seeds','Annotate outcomes, artifacts,\nand collateral changes.')]:
    box(s,x,2.65,3.55,2.65,'#FFFFFF',LINE)
    txt(s,n,x+.18,2.82,.9,.62,size=28,color=TEAL,bold=True)
    txt(s,title,x+.18,3.55,3.18,.58,size=18,bold=True)
    txt(s,body,x+.18,4.28,3.18,.95,size=16)
box(s,.85,5.8,11.55,.83,PALE)
txt(s,'Measure Δ unwanted-outcome rate by group + artifact cost + uncertainty.',1.05,5.99,11.1,.5,size=20,color=TEAL,bold=True)

s=slide('The experiment is reproducible and reviewable','Scripts, grids, metrics, raw activation maps, and a blank annotation sheet are saved.','Artifacts',
    'The new bias experiment script has an analysis-only mode and a complete replay mode. Outputs include the dataset split, candidate associations, test activation metrics, intervention pixel changes, per-image raw maps, individual PNGs, 36 grids, and a markdown report. The annotation sheet has 432 rows, one per intervention, with blank fields for outcomes and preference; those blanks must not be mistaken for results. Existing notebooks and baseline images were preserved. The separate generic SAE training script remains available but its prompt split is not suitable as a balanced portrait audit without adaptation.',
    'sdxl-unbox/scripts/sae_bias_experiment.py; SAE_BIAS_EXPERIMENT.md; output REPORT.md')
txt(s,'Run analysis only',.95,2.55,4.5,.45,size=20,bold=True,color=TEAL)
box(s,.95,3.13,11.4,1.18,PALE)
txt(s,'python -I -u sdxl-unbox/scripts/sae_bias_experiment.py\n  --output baseline/new_bias_analysis --analysis-only',1.14,3.3,11,.88,size=18)
txt(s,'Run all-category sweeps',.95,4.65,5.3,.45,size=20,bold=True,color=TEAL)
box(s,.95,5.22,11.4,1.18,PALE)
txt(s,'python -I -u sdxl-unbox/scripts/sae_bias_experiment.py\n  --output baseline/new_bias_interventions --seeds 40 41 42',1.14,5.39,11,.88,size=18)

s=slide('Appendix · the initial SAE training pilot','A small new dictionary was trained, then excluded from the racial-bias experiment.','Training provenance',
    'Before the portrait scope was clarified, a separate Top-K SAE was trained on ten varied prompts with four seeds per prompt, with two additional prompts held out. The actual dataset has 10,240 training vectors and 2,048 validation vectors; the selected block has 256 spatial positions per image, below the collection cap of 512. The new dictionary has 5,120 features and k=20 and was trained for 500 Adam steps at learning rate 1e-4. Held-out normalized MSE fell from 1.221 to 0.976, near a held-out-mean prediction baseline. 86.6 percent of features were unused on that validation sample, not globally dead. This pilot demonstrates trainability, not adequate feature quality, and its IDs do not transfer to the pretrained dictionary.',
    'baseline/sae_experiment_20261009/RESULTS.md; training_metrics.json; sae_experiment.py')
pilot=json.loads((ROOT/'baseline/sae_experiment_20261009/training_metrics.json').read_text())
card(s,.9,2.6,5.6,2.65,'Training configuration','1,280 → 5,120 dimensions; k=20\n10,240 train / 2,048 held-out vectors\n500 Adam steps; learning rate 10⁻⁴')
card(s,6.75,2.6,5.6,2.65,'Observed fit','Training batch MSE: 7.774 → 0.856\nHeld-out normalized MSE: 1.221 → 0.976\n86.6% unused on held-out sample',ORANGE)
txt(s,'For the portrait audit, reuse of the original pretrained k=10 SAE preserves old feature IDs.',.95,5.85,11.5,.8,size=20,bold=True,color=TEAL)

s=slide('Appendix · a normalization issue in the old notebook','The new analysis leaves old artifacts intact and reports quantities with correct denominators.','Measurement audit',
    'The old race-prevalence notebook sums active spatial positions for each feature and divides by the number of generated images. Its comment calls the result a fraction of positions, but the value is the mean number of active positions per image. A true spatial active fraction divides by the total number of positions, or averages a per-image fraction when shapes are constant. The new cached-data association uses spatial maxima because that is what the existing feature file contains, and labels it accordingly. The new replay metrics calculate spatial means, maxima, and actual active fractions. No historical JSON, CSV, or notebook output was overwritten.',
    'baseline/sae_race_feature_analysis.ipynb; existing_feature_138_527.csv; test_activation_metrics.csv')
card(s,.9,2.65,5.65,2.95,'Old “freq” field','Σ active spatial positions / number of images\n\nMeaning: active positions per image\nCan exceed 1.',ORANGE)
card(s,6.75,2.65,5.65,2.95,'New replay fraction','mean over positions [activation > 0]\n\nMeaning: spatial fraction per image\nBounded between 0 and 1.')
txt(s,'Cached NPZ remains spatial-max pooled; its image-active fraction is a different statistic.',.95,6.05,11.5,.55,size=17,color=MUTED)

s=slide('Appendix · local evidence and exact configuration','All reported quantitative results are derived from workspace artifacts.','Evidence index',
    'The slides are grounded in local output files, rather than paper claims or external benchmarks. The pretrained checkpoint config lists 1,280 input dimensions, 5,120 features, k=10, auxk=256, and a dead-steps threshold of 2,441. The run uses cached SDXL Turbo and the fp16-fixed VAE, a single GPU, one 512-square denoising step, and the settings shown here. The working Python environment was isolated with -I to avoid a broken user-level torch import. Model startup initially waited on filesystem reads; after loading, generation and intervention checks completed. These environment details explain execution but are not research results.',
    'FINDINGS.md; settings.json; intervention_metrics.csv; pretrained final/config.json; build_sae_presentation.py')
entries=[('Findings','baseline/sae_bias_interventions_20261009/FINDINGS.md'),('Method','sdxl-unbox/scripts/sae_bias_experiment.py'),('Associations','feature_associations.csv; cached_data_split.npz'),('Controls','intervention_metrics.csv; settings.json'),('Visuals / review','REPORT.md; */feature*_comparison.png; human_review.csv'),('Training pilot','baseline/sae_experiment_20261009/RESULTS.md')]
for i,(label,path) in enumerate(entries):
    y=2.55+i*.55
    txt(s,label,.95,y,2.45,.4,size=15,bold=True,color=TEAL)
    txt(s,path,3.45,y,8.9,.44,size=13,color=MUTED)
txt(s,'Interpretation: successful controlled edits; no demonstrated fairness or quality improvement.',.95,6.3,11.5,.5,size=17,bold=True)

PPT = OUT / 'SAE_Racial_Bias_Audit.pptx'
prs.save(PPT)
for i,im in enumerate(previews,1):
    im.save(OUT/'previews'/f'slide_{i:02d}.png')
thumb_w, thumb_h = 400, 225
contact = Image.new('RGB',(thumb_w*4, (math.ceil(len(previews)/4))*(thumb_h+26)),'#DCE3E9')
draw=ImageDraw.Draw(contact)
for i,im in enumerate(previews):
    x,y=(i%4)*thumb_w,(i//4)*(thumb_h+26)
    contact.paste(im.resize((thumb_w,thumb_h),Image.Resampling.LANCZOS),(x,y+26))
    draw.text((x+8,y+4),f'Slide {i+1:02d}',font=font(10),fill=INK)
contact.save(OUT/'slide_overview.png')
speaker=['# Speaker notes — SAE racial-bias audit','', '20 slides; approximately 15–20 minutes.','',
         'The PNG previews visualize the authored layout; they are not a PowerPoint-engine render.','']
for i,(title,note,source) in enumerate(notes,1):
    speaker.extend([f'## Slide {i}: {title}','',note,'',f'Local evidence: `{source}`',''])
(OUT/'Speaker_Notes.md').write_text('\n'.join(speaker))

# Verify the archive, XML, notes, slide geometry, and observed control counts.
with ZipFile(PPT) as archive:
    assert archive.testzip() is None
    for name in archive.namelist():
        if name.endswith('.xml'):
            ET.fromstring(archive.read(name))
    assert len([n for n in archive.namelist() if n.startswith('ppt/slides/slide') and n.endswith('.xml')])==20
    assert len([n for n in archive.namelist() if n.startswith('ppt/notesSlides/notesSlide') and n.endswith('.xml')])==20
check=Presentation(PPT)
for i,s in enumerate(check.slides,1):
    for sh in s.shapes:
        assert sh.left >= 0 and sh.top >= 0
        assert sh.left+sh.width <= check.slide_width+Inches(.03), (i,sh.name)
        assert sh.top+sh.height <= check.slide_height+Inches(.03), (i,sh.name)
assert len(metrics)==432
assert sum(float(row['strength'])==1 for row in metrics)==108
assert max(float(row['saved_baseline_replay_error']) for row in metrics)==0
assert max(float(row['pixel_change']) for row in metrics if float(row['strength'])==1)==0
print(json.dumps({'pptx':str(PPT),'slides':len(check.slides),'notes':len(notes),'size_mb':round(PPT.stat().st_size/1024**2,2),'validation':'ZIP, XML, notes, slide bounds, text layout, and reported counts passed'},indent=2))
