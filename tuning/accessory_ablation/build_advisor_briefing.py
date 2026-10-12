"""Build a plain monochrome advisor briefing for the SAE ablation pilots."""
import sys
from pathlib import Path
sys.path.insert(0, '/tmp/sae_ppt_dependencies')
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

ROOT = Path(__file__).resolve().parents[2]
OUT = ROOT / 'tuning/accessory_ablation/advisor_briefing.pptx'
BLACK = RGBColor(0, 0, 0)
WHITE = RGBColor(255, 255, 255)
prs = Presentation()
prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
prs.core_properties.title = 'SAE Ablation Pilot: Why Earring and Headscarf Removal Was Not Reliable'
prs.core_properties.subject = 'Advisor briefing on the feature ablation and C2C box-masked follow-up'
prs.core_properties.author = 'Research experiment'
blank = prs.slide_layouts[6]


def add_text(slide, text, x, y, w, h, size=18, bold=False, align=None):
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = shape.text_frame
    tf.clear(); tf.word_wrap = True
    tf.margin_left = Inches(.02); tf.margin_right = Inches(.02)
    tf.margin_top = Inches(.01); tf.margin_bottom = Inches(.01)
    p = tf.paragraphs[0]; p.text = text
    p.font.name = 'Arial'; p.font.size = Pt(size); p.font.bold = bold; p.font.color.rgb = BLACK
    if align is not None: p.alignment = align
    return shape


def add_slide(title, subtitle, notes):
    s = prs.slides.add_slide(blank)
    s.background.fill.solid(); s.background.fill.fore_color.rgb = WHITE
    add_text(s, title, .65, .38, 12.05, .58, 26, True)
    if subtitle: add_text(s, subtitle, .68, 1.02, 11.95, .42, 13)
    line = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(.68), Inches(1.53), Inches(11.95), Inches(.012))
    line.fill.solid(); line.fill.fore_color.rgb = BLACK; line.line.fill.background()
    add_text(s, f'{len(prs.slides):02d}', 12.15, 7.12, .45, .2, 9, False, PP_ALIGN.RIGHT)
    s.notes_slide.notes_text_frame.text = notes
    return s


def body(s, text, x, y, w, h, size=18, bold=False):
    return add_text(s, text, x, y, w, h, size, bold)


def table(s, x, y, widths, rows, row_h=.42, font_size=14):
    shape = s.shapes.add_table(len(rows), len(widths), Inches(x), Inches(y),
                               Inches(sum(widths)), Inches(row_h*len(rows)))
    t = shape.table
    for i, width in enumerate(widths): t.columns[i].width = Inches(width)
    for ri, row in enumerate(rows):
        for ci, value in enumerate(row):
            cell = t.cell(ri, ci); cell.text = str(value)
            cell.fill.solid(); cell.fill.fore_color.rgb = WHITE
            cell.margin_left = Inches(.07); cell.margin_right = Inches(.05)
            cell.margin_top = Inches(.025); cell.margin_bottom = Inches(.02)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            for p in cell.text_frame.paragraphs:
                p.font.name='Arial'; p.font.size=Pt(font_size); p.font.bold=(ri==0)
                p.font.color.rgb=BLACK
            # Explicit thin black cell borders in the underlying XML.
            tcPr = cell._tc.get_or_add_tcPr()
            from pptx.oxml.xmlchemy import OxmlElement
            from pptx.oxml.ns import qn
            for edge in ('a:lnL','a:lnR','a:lnT','a:lnB'):
                ln=OxmlElement(edge); ln.set('w','6350')
                sf=OxmlElement('a:solidFill'); c=OxmlElement('a:srgbClr'); c.set('val','000000'); sf.append(c); ln.append(sf)
                dash=OxmlElement('a:prstDash'); dash.set('val','solid'); ln.append(dash)
                tcPr.append(ln)
    return shape

s=add_slide('SAE ablation did not reliably remove earrings or headscarves',
            'Advisor briefing · matched-seed pilot followed by a C2C-box-masked earring test',
            'Suggested opening: We ran a controlled feature ablation on the same prompts and seeds. The replay and no-op checks worked. The selected SAE features changed pixels, but visual review did not show reliable removal. A later box-masked test localized the edits and reduced collateral image change, while the earrings remained visible. The evidence points to weak feature-to-concept alignment and an evaluation that measured image change rather than object removal.')
body(s,'What worked: exact replay and controlled spatial intervention.',.9,2.25,11.5,.65,23)
body(s,'What did not: consistent control of the target object.',.9,3.1,11.5,.65,23)
body(s,'Main explanation: feature 4858 was a weak held-out earring indicator, and removing its contribution did not remove the earrings in the inspected images.',.9,4.35,11.3,1.15,19)
body(s,'Conclusion is limited to this small exploratory pilot; it does not establish a fairness effect.',.9,6.1,11.4,.5,15)

s=add_slide('What we did in the first pilot',
            'Same prompt and seed isolate the feature edit from ordinary generation variation.',
            'The first pilot replayed 16 same-prompt validation images using their original seeds. There were eight examples per concept: four detector-positive and four detector-negative according to unreviewed Florence-2 labels. Candidate IDs were chosen on discovery data: headscarf 1257 and earrings 4858. The diffusion model and pretrained SAE were frozen. At the tapped U-Net block, we scaled one SAE feature contribution while preserving the rest of the residual output. Strength 1 leaves the image unchanged; 0 removes that contribution; 0.5 keeps half; negative values reverse and enlarge the change relative to the original contribution. Matched random features were 2539 for headscarf and 2940 for earrings. The baseline replay was pixel-identical for all 16 images, and strength 1 was an exact no-op.')
body(s,'16 validation images total · 8 per concept · 4 automatic-label positives and 4 negatives per concept',.82,1.9,11.7,.55,18)
body(s,'Frozen pretrained SAE and diffusion model · one U-Net feature contribution scaled per run',.82,2.65,11.7,.55,18)
table(s,.9,3.55,[2.2,2.4,2.4,2.6], [['Concept','Candidate','Random control','Held-out AUC'],['Headscarf','1257','2539','0.228'],['Earrings','4858','2940','0.407']],.52,16)
body(s,'AUC values below 0.5 indicate these candidate activations ranked held-out positives below negatives in this comparison.',.9,5.55,11.4,.75,16)
body(s,'Strengths tested: −1, −0.5, 0, 0.5, 1. Pixel MAE records visual change, not target presence.',.9,6.45,11.6,.35,14)

s=add_slide('The first pilot showed change, not reliable object control',
            'Negative strength perturbed more; the random earring feature perturbed the image more than the candidate.',
            'Explain this slide carefully: The reported metric is mean absolute RGB pixel difference from the unchanged replay on a 0–255 scale. It answers how much the image changed, not whether an earring or headscarf disappeared. At strength −1, headscarf candidate change averaged 1.127 versus 0.761 for its random control. For earrings, the candidate averaged 6.334 while the random feature averaged 10.386. In inspected examples, neither zeroing nor negative scaling clearly removed the target. The low AUCs reinforce that these particular candidate IDs were not reliable semantic detectors on held-out data.')
table(s,.9,2.05,[2.0,1.5,2.4,2.5,2.4], [['Concept','Strength','Candidate MAE','Random MAE','Visual result'],['Headscarf','0','0.765','0.577','No clear removal'],['Headscarf','−1','1.127','0.761','No clear removal'],['Earrings','0','3.658','6.359','No clear removal'],['Earrings','−1','6.334','10.386','No clear removal']],.54,14)
body(s,'MAE is RGB pixel change on a 0–255 scale. More change is not evidence of better removal.',.9,5.25,11.3,.52,17)
body(s,'Small sample; source labels were unreviewed Florence-2 outputs, and detector negatives may include missed objects.',.9,5.95,11.5,.78,16)

s=add_slide('C2C boxes localized feature 4858 but earrings remained',
            'Follow-up: four same-prompt validation positives, with intervention restricted to C2C earring boxes.',
            'The follow-up mapped C2C boxes from 512 by 512 image coordinates to the SAE spatial code grid and expanded the mask by one grid cell. It compared feature 4858 reduced globally, feature 4858 reduced only inside boxes, and random feature 2940 reduced inside the same boxes. There were four examples. Strength 0 removed the contribution; 0.5 retained half; 1 was the no-op. Box masking reduced collateral image change: global candidate whole-image MAE was 4.354, versus 2.799 for box-only candidate. Inside-box MAE was 9.612 globally and 6.092 for box-only. The earrings remained in visual review. The random control also changed some images, but was zero for some because that feature was inactive at the masked locations; it is not a fully matched effective perturbation control. So the localization implementation works, while the chosen feature still does not provide reliable object removal.')
table(s,.72,1.9,[2.35,2.5,2.15,2.15,2.15], [['Condition at strength 0','Feature','Whole-image MAE','Inside-box MAE','Outside-box MAE'],['Candidate, global','4858','4.354','9.612','4.318'],['Candidate, C2C boxes','4858','2.799','6.092','2.775'],['Random, C2C boxes','2940','0.831','1.578','0.826']],.56,13)
body(s,'Four images · source replay exact · strength 1 exact no-op · one-cell spatial padding',.85,4.45,11.7,.45,16)
body(s,'Interpretation: box masking reduced the spatial footprint, but did not make feature 4858 an earring-removal control.',.85,5.1,11.65,.78,18)
body(s,'The random feature was inactive in some boxes, so improve the control by matching on in-box activation and perturbation magnitude.',.85,6.15,11.65,.55,15)

s=add_slide('Why reliability is still low',
            'The bottleneck is target-feature validity and outcome measurement, not replay fidelity.',
            'Three limitations explain the result. First, candidate features were selected from discovery data but had weak held-out association: AUC 0.228 and 0.407. Second, pixel MAE is not an earring-presence measure; large changes can affect hair, skin, or background while leaving the jewelry intact. Third, sample sizes were only 16 in the initial pilot and four in the box-masked follow-up. Automatic labels and C2C boxes were not visually audited in this experiment. Also, a spatial box only says where to apply the edit; it cannot make a semantically mixed or unrelated feature specific to earrings. The random control should be matched for activation inside the same boxes, because a globally matched feature may be inactive locally.')
body(s,'1. Weak feature validity: held-out AUC was 0.228 for headscarf and 0.407 for earrings.',.9,1.95,11.4,.7,18)
body(s,'2. Indirect outcome: image MAE measures change, not whether the earring is still present.',.9,2.9,11.4,.7,18)
body(s,'3. Limited review: 16 initial images; 4 box-masked images; labels and boxes need human checks.',.9,3.85,11.4,.7,18)
body(s,'4. Weak local control comparison: random feature 2940 is inactive in some intervention boxes.',.9,4.8,11.4,.7,18)
body(s,'A spatial mask narrows where the edit is applied; it cannot repair poor feature-to-concept alignment.',.9,6.05,11.4,.6,19,True)

s=add_slide('How to make the next test more reliable',
            'Separate feature validation, controlled editing, and object-removal measurement.',
            'Suggested plan: First, build a human-verified set of earring boxes and object-presence labels; include multiple prompt groups, seeds, and earring styles. On a held-out split, test which SAE features are active in those boxes and predict human labels. Keep feature selection separate from evaluation. If no single feature validates, test a small, prespecified feature set or a different representation rather than repeatedly selecting on the test set. Second, compare candidate edits against no-op and random controls matched for in-box activation and pixel-change scale, and tune strength only on a development split. Third, measure object presence with blinded human review or an independently validated detector, along with collateral changes outside the box and image quality. Report per-image outcomes and uncertainty, not only aggregate MAE. Set a success rule in advance: a clear decrease in earring presence, robust across seeds, with low unrelated change.')
body(s,'1 · Validate the feature on a human-checked, held-out earring set; report in-box precision/recall or AUC.',.82,1.9,11.8,.72,17)
body(s,'2 · Use more seeds and prompt groups; manually verify boxes and positive/negative labels.',.82,2.8,11.8,.65,17)
body(s,'3 · Match random controls on local activation and perturbation size; tune edit strength only on development data.',.82,3.65,11.8,.72,17)
body(s,'4 · Score earring presence directly with blinded human review or an independently checked detector.',.82,4.55,11.8,.65,17)
body(s,'5 · Track collateral change outside the box and image quality; publish per-image results and uncertainty.',.82,5.4,11.8,.65,17)
body(s,'Success criterion: target presence drops consistently across held-out seeds while unrelated regions remain stable.',.82,6.25,11.8,.55,17,True)

s=add_slide('Reproducibility and files',
            'The saved reports contain the grids, per-image outputs, and measurements.',
            'Close by giving the advisor the two reports and their scripts. The initial experiment script prepares its manifest and runs the replayed feature interventions. The box-masked script reproduces the spatial mask comparison. Reports and CSVs are in the same folders. The runs verified exact source replay and exact no-op at strength one. State clearly that this is a small exploratory result and that the next experiment should test candidate validity and object presence directly.')
body(s,'Initial ablation report: tuning/accessory_ablation/REPORT.md',.9,2.05,11.4,.48,18)
body(s,'Initial experiment script: tuning/accessory_ablation.py',.9,2.75,11.4,.48,18)
body(s,'Box-masked follow-up report: tuning/box_masked_earring/REPORT.md',.9,3.45,11.4,.48,18)
body(s,'Box-masked follow-up script: tuning/box_masked_earring.py',.9,4.15,11.4,.48,18)
body(s,'Metrics: tuning/accessory_ablation/metrics.csv and tuning/box_masked_earring/metrics.csv',.9,4.85,11.6,.72,18)
body(s,'Advisor summary: the intervention was technically reproducible and spatially localizable; reliable semantic removal was not demonstrated.',.9,6.05,11.5,.7,19,True)

OUT.parent.mkdir(exist_ok=True)
prs.save(OUT)
print(OUT)
